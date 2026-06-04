"""Best-effort text extraction from attachment bytes.

Supported: PDF (PyMuPDF), DOCX (python-docx), XLSX (openpyxl), and plain text.
Anything else returns ``None`` so only metadata is indexed. Extraction never
raises: failures are logged and degrade to ``None``.
"""

from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)

_PLAINTEXT_MIME_PREFIXES = ("text/",)


def extract_text(filename: str | None, mime_type: str | None, data: bytes) -> str | None:
    name = (filename or "").lower()
    mime = (mime_type or "").lower()
    try:
        if name.endswith(".pdf") or "pdf" in mime:
            return _pdf(data)
        if name.endswith(".docx") or "wordprocessingml" in mime:
            return _docx(data)
        if name.endswith(".xlsx") or "spreadsheetml" in mime:
            return _xlsx(data)
        if name.endswith((".txt", ".csv", ".md", ".log")) or mime.startswith(_PLAINTEXT_MIME_PREFIXES):
            return _plain(data)
    except Exception as exc:  # noqa: BLE001 - extraction must never break sync
        logger.warning("Extraction failed for %s (%s): %s", filename, mime_type, exc)
    return None


def _clean(text: str) -> str | None:
    text = " ".join(text.split())
    return text or None


def _pdf(data: bytes) -> str | None:
    import fitz  # PyMuPDF

    parts: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            parts.append(page.get_text())
    return _clean("\n".join(parts))


def _docx(data: bytes) -> str | None:
    import docx

    document = docx.Document(io.BytesIO(data))
    return _clean("\n".join(p.text for p in document.paragraphs))


def _xlsx(data: bytes) -> str | None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            parts.append(" ".join(str(c) for c in row if c is not None))
    wb.close()
    return _clean("\n".join(parts))


def _plain(data: bytes) -> str | None:
    return _clean(data.decode("utf-8", errors="replace"))
